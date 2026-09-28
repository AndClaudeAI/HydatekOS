"""Build PowerPoint files for Hyda Slides' import tests.

  python-made.pptx   written by python-pptx on its default (4:3) template
  office-made.pptx   the same deck re-saved by LibreOffice Impress
  office-features.pptx  Hyda Slides' own export of its feature deck (tables,
                     charts, shapes, arrows, animations, footers), re-saved by
                     LibreOffice Impress (made with the host tests' export)

These tools are test tools on the build machine only.
Run: python3 make.py   (needs python-pptx, Pillow and soffice)
"""
import os, subprocess, shutil, tempfile
from pptx import Presentation
from pptx.util import Pt, Emu, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))

def main():
    prs = Presentation()
    # 1: title slide
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Quarterly review"
    s.placeholders[1].text = "Lagos office, September 2026"
    s.notes_slide.notes_text_frame.text = "Welcome everyone.\nIntroduce the team."
    # 2: title and content with levels and formatting
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "What we shipped"
    tf = s.placeholders[1].text_frame
    tf.text = "Solar mini-grids"
    p = tf.add_paragraph(); p.text = "Kano markets"; p.level = 1
    p = tf.add_paragraph()
    r = p.add_run(); r.text = "Bold "; r.font.bold = True
    r = p.add_run(); r.text = "and italic"; r.font.italic = True
    p = tf.add_paragraph(); p.text = "Mobile money"
    # 3: blank with shapes, a text box, a picture and a group
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1.5))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xC0, 0x62, 0x2B)
    box.text_frame.text = "Revenue up 12%"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    ov = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5), Inches(1), Inches(2), Inches(2))
    ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0x2F, 0x6F, 0xEB)
    tb = s.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    tb.text_frame.text = "A plain text box"
    tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    img = os.path.join(tempfile.mkdtemp(), "pic.png")
    im = Image.new("RGB", (64, 32), (0x1E, 0x1B, 0x2C))
    for x in range(32):
        for y in range(32):
            im.putpixel((x, y), (0xF2, 0xB5, 0x44))
    im.save(img)
    s.shapes.add_picture(img, Inches(6), Inches(4), Inches(2), Inches(1))
    grp = s.shapes.add_group_shape()
    g1 = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(5), Inches(1), Inches(1))
    g1.fill.solid(); g1.fill.fore_color.rgb = RGBColor(0x0E, 0x5A, 0x43)
    g2 = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(5), Inches(1), Inches(1))
    g2.fill.solid(); g2.fill.fore_color.rgb = RGBColor(0xF5, 0xC5, 0x42)
    # 4: a table, charts, an arrow, a turned star and a gradient
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Tables and charts"
    tbl = s.shapes.add_table(3, 3, Inches(0.5), Inches(1.5), Inches(4), Inches(1.2)).table
    for r, row in enumerate([["", "Q1", "Q2"], ["North", "5", "6"], ["South", "8", "7"]]):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    cd = CategoryChartData()
    cd.categories = ["East", "West", "Midwest"]
    cd.add_series("Q1 Sales", (19.2, 21.4, 16.7))
    cd.add_series("Q2 Sales", (22.3, 28.6, 15.2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.5), Inches(4.5), Inches(2.5), cd)
    pd = CategoryChartData()
    pd.categories = ["A", "B", "C"]
    pd.add_series("Share", (50, 30, 20))
    s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(0.5), Inches(3), Inches(3), Inches(2.5), pd)
    s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(3.5), Inches(4.2), Inches(3), Inches(2.5), cd)
    cx = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6.5), Inches(5), Inches(9), Inches(6))
    ln = cx.line._get_or_add_ln()
    ln.append(etree.SubElement(ln, qn("a:tailEnd"), {"type": "triangle"}))
    st = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(8), Inches(4), Inches(1.5), Inches(1.5))
    st.rotation = 45
    gr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(6.2), Inches(2.5), Inches(1))
    gr.fill.gradient()
    gr.fill.gradient_stops[0].color.rgb = RGBColor(0xFF, 0, 0)
    gr.fill.gradient_stops[1].color.rgb = RGBColor(0, 0, 0xFF)
    out = os.path.join(HERE, "python-made.pptx")
    prs.save(out)
    # the same deck, re-saved by LibreOffice
    tmp = tempfile.mkdtemp()
    shutil.copy(out, os.path.join(tmp, "in.pptx"))
    subprocess.run(["soffice", "--headless", "--convert-to", "odp", "--outdir", tmp, os.path.join(tmp, "in.pptx")], check=True, capture_output=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pptx", "--outdir", os.path.join(tmp, "o"), os.path.join(tmp, "in.odp")], check=True, capture_output=True)
    shutil.copy(os.path.join(tmp, "o", "in.pptx"), os.path.join(HERE, "office-made.pptx"))
    # Hyda Slides' feature deck, through LibreOffice
    tmp = tempfile.mkdtemp()
    feat = os.path.join(tmp, "features.pptx")
    subprocess.run(["cargo", "test", "--release", "--quiet", "export_features", "--", "--ignored"], cwd=os.path.join(HERE, "..", ".."), env=dict(os.environ, HYDA_OUT=feat), check=True, capture_output=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "odp", "--outdir", tmp, feat], check=True, capture_output=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pptx", "--outdir", os.path.join(tmp, "o"), os.path.join(tmp, "features.odp")], check=True, capture_output=True)
    shutil.copy(os.path.join(tmp, "o", "features.pptx"), os.path.join(HERE, "office-features.pptx"))

if __name__ == "__main__":
    main()
